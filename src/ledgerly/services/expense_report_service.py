import io
import re
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExpenseReportGenerator:

    def __init__(self, df: pd.DataFrame):
        self.df = df.copy() if df is not None else pd.DataFrame()

        # Data cleanup & standard types
        if not self.df.empty:
            if "posted_date" in self.df.columns:
                self.df["posted_date"] = pd.to_datetime(self.df["posted_date"])
            if "amount" in self.df.columns:
                self.df["amount"] = pd.to_numeric(self.df["amount"], errors="coerce").fillna(0.0)

        # Styling Palette
        self.navy_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        self.soft_blue_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        self.kpi_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        self.accent_fill = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")

        self.title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
        self.section_font = Font(name="Calibri", size=13, bold=True, color="1F4E78")
        self.header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        self.sub_header_font = Font(name="Calibri", size=10, bold=True, color="1F4E78")
        self.bold_font = Font(name="Calibri", size=10, bold=True)
        self.regular_font = Font(name="Calibri", size=10)

        self.thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        self.double_bottom_border = Border(
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="double", color="1F4E78"),
        )

    def generate_report(
        self,
        months_selected: list[str],
        user_info: dict = None,
        master_file_bytes: bytes = None,
    ) -> io.BytesIO:
        """Generates or updates the Master Expense Report workbook."""
        if master_file_bytes:
            wb = openpyxl.load_workbook(io.BytesIO(master_file_bytes))
            if "Overview" in wb.sheetnames:  # Migrate legacy 'Overview' tab
                wb.remove(wb["Overview"])
        else:
            wb = openpyxl.Workbook()
            wb.remove(wb.active)

        # 1. Generate/Overwrite selected Monthly Tabs
        for month_str in months_selected:
            if month_str in wb.sheetnames:
                wb.remove(wb[month_str])

            month_df = self._filter_month_data(month_str)
            ws = wb.create_sheet(title=month_str)
            self._build_monthly_sheet(ws, month_df, month_str)

        # 2. Discover all month tabs in workbook (chronological ascending)
        date_pattern = re.compile(r"^\d{4}-\d{2}$")
        all_month_sheets = sorted(
            [s for s in wb.sheetnames if date_pattern.match(s)],
            key=lambda x: datetime.strptime(x, "%Y-%m"),
        )

        years_present = sorted(list(set(m.split("-")[0] for m in all_month_sheets)))

        # 3. Generate/Rebuild Yearly Overviews
        for year in years_present:
            year_title = f"{year} Overview"
            if year_title in wb.sheetnames:
                wb.remove(wb[year_title])
            
            ws_year = wb.create_sheet(title=year_title)
            months_in_year = [m for m in all_month_sheets if m.startswith(year)]
            self._build_yearly_overview_sheet(ws_year, year, months_in_year, all_month_sheets)

        # 4. Generate/Rebuild Master Overview
        if "Master Overview" in wb.sheetnames:
            wb.remove(wb["Master Overview"])
        
        ws_master = wb.create_sheet(title="Master Overview")
        self._build_master_overview_sheet(ws_master, years_present, all_month_sheets)

        # 5. Reorder sheet tabs only: Master Overview -> Yearly Overviews (descending) -> Monthly Tabs (descending)
        self._sort_sheets_chronologically(wb)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def _get_year_status_label(self, year_str: str, month_sheets: list[str]) -> str:
        """Determines year status: '2026 (Current)', '2025 (Missing: Nov, Dec)', or '2024'."""
        current_year = datetime.now().year
        try:
            year_int = int(year_str)
        except ValueError:
            return year_str

        present_months = sorted([
            int(m.split("-")[1]) for m in month_sheets if m.startswith(year_str)
        ])
        all_months = list(range(1, 13))
        missing_months = [m for m in all_months if m not in present_months]

        month_names = {
            1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "May", 6: "Jun",
            7: "Jul", 8: "Aug", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dec"
        }

        if year_int == current_year:
            return f"{year_str} (Current)"
        elif year_int < current_year:
            if missing_months:
                missing_str = ", ".join(month_names[m] for m in missing_months)
                return f"{year_str} (Missing: {missing_str})"
            return year_str
        else:
            return f"{year_str} (Future)"

    def _filter_month_data(self, month_str: str) -> pd.DataFrame:
        if self.df.empty or "posted_date" not in self.df.columns:
            return pd.DataFrame()
        temp_df = self.df.copy()
        temp_df["month_key"] = temp_df["posted_date"].dt.strftime("%Y-%m")
        return temp_df[temp_df["month_key"] == month_str].sort_values("posted_date")

    def _is_transfer(self, row) -> bool:
        cat = str(row.get("category", "")).lower()
        subcat = str(row.get("subcategory", "")).lower()
        return cat == "financial & transfers" and subcat in [
            "account transfer",
            "credit card payment",
            "transfer",
            "loan proceeds",
        ]

    def _sort_sheets_chronologically(self, wb: openpyxl.Workbook):
        """Reorders sheet tabs: Master Overview -> YYYY Overviews (descending) -> Monthly Tabs (descending)."""
        date_pattern = re.compile(r"^\d{4}-\d{2}$")
        month_sheets = sorted(
            [s for s in wb.sheetnames if date_pattern.match(s)],
            key=lambda x: datetime.strptime(x, "%Y-%m"),
            reverse=True,
        )
        years = sorted(list(set(m.split("-")[0] for m in month_sheets)), reverse=True)

        target_order = []

        # 1. Master Overview first
        if "Master Overview" in wb.sheetnames:
            target_order.append("Master Overview")

        # 2. All Yearly Overview tabs grouped upfront in reverse chronological order
        for y in years:
            y_ov = f"{y} Overview"
            if y_ov in wb.sheetnames:
                target_order.append(y_ov)

        # 3. All Monthly Tabs in reverse chronological sequence
        target_order.extend(month_sheets)

        # 4. Append remaining unmapped sheets
        leftovers = [s for s in wb.sheetnames if s not in target_order]
        target_order.extend(leftovers)

        wb._sheets = [wb[title] for title in target_order]

    # =========================================================================
    # MONTHLY SHEET BUILDER
    # =========================================================================
    def _build_monthly_sheet(self, ws, month_df: pd.DataFrame, month_str: str):
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = f"Ledgerly Statement — {month_str}"
        ws["A1"].font = self.title_font

        # Calculate monthly totals for header banner
        income, expenses, net, count = 0.0, 0.0, 0.0, 0
        if not month_df.empty:
            kpi_df = month_df[~month_df.apply(self._is_transfer, axis=1)]
            income = kpi_df[kpi_df["amount"] > 0]["amount"].sum()
            expenses = kpi_df[kpi_df["amount"] < 0]["amount"].sum()
            net = income + expenses
            count = len(month_df)

        headers = ["Total Income", "Total Expenses", "Net Cashflow", "Total Txns"]
        vals = [income, abs(expenses), net, count]

        for i, (h, v) in enumerate(zip(headers, vals), start=1):
            cell_h = ws.cell(row=3, column=i, value=h)
            cell_v = ws.cell(row=4, column=i, value=v)
            cell_h.fill = self.soft_blue_fill
            cell_h.font = self.sub_header_font
            cell_v.fill = self.kpi_fill
            cell_v.font = self.bold_font
            if i in [1, 2, 3]:
                cell_v.number_format = '"$"#,##0.00'
            else:
                cell_v.number_format = '#,##0'

        # Transaction Ledger Table
        start_row = 7
        cols = ["Date", "Description", "Category", "Subcategory", "Amount"]
        for col_idx, col_name in enumerate(cols, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=col_name)
            cell.fill = self.navy_fill
            cell.font = self.header_font

        curr_row = start_row + 1
        if not month_df.empty:
            for _, row in month_df.iterrows():
                ws.cell(row=curr_row, column=1, value=row["posted_date"].strftime("%Y-%m-%d")).font = self.regular_font
                ws.cell(row=curr_row, column=2, value=str(row.get("description", ""))).font = self.regular_font
                ws.cell(row=curr_row, column=3, value=str(row.get("category", ""))).font = self.regular_font
                ws.cell(row=curr_row, column=4, value=str(row.get("subcategory", ""))).font = self.regular_font

                amt_cell = ws.cell(row=curr_row, column=5, value=float(row["amount"]))
                amt_cell.font = self.regular_font
                amt_cell.number_format = '"$"#,##0.00;("$"#,##0.00);"-"'
                curr_row += 1

        # Summary Total Row
        ws.cell(row=curr_row, column=4, value="Total").font = self.bold_font
        total_cell = ws.cell(row=curr_row, column=5, value=f"=SUM(E{start_row+1}:E{max(curr_row-1, start_row+1)})")
        total_cell.font = self.bold_font
        total_cell.number_format = '"$"#,##0.00;("$"#,##0.00);"-"'
        total_cell.border = self.double_bottom_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # =========================================================================
    # YEARLY OVERVIEW BUILDER
    # =========================================================================
    def _build_yearly_overview_sheet(self, ws, year_str: str, month_sheets: list[str], all_month_sheets: list[str]):
        ws.views.sheetView[0].showGridLines = True
        
        status_label = self._get_year_status_label(year_str, all_month_sheets)
        ws["A1"] = f"{status_label} Financial Overview" if "(" in status_label else f"Financial Overview {year_str}"
        ws["A1"].font = self.title_font

        # BLOCK 2: Monthly Cashflow Breakdown
        curr_row = 7
        ws.cell(row=curr_row, column=1, value="Monthly Cashflow Breakdown").font = self.section_font
        curr_row += 1

        headers = ["Month", "Total Income", "Total Expenses", "Net Cashflow", "Savings Rate", "Txn Count"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.fill = self.navy_fill
            cell.font = self.header_font

        start_breakdown = curr_row + 1
        curr_row += 1

        for sheet_name in month_sheets:
            ws.cell(row=curr_row, column=1, value=sheet_name).font = self.regular_font
            
            inc_cell = ws.cell(row=curr_row, column=2, value=f"='{sheet_name}'!A4")
            exp_cell = ws.cell(row=curr_row, column=3, value=f"='{sheet_name}'!B4")
            net_cell = ws.cell(row=curr_row, column=4, value=f"=B{curr_row}-C{curr_row}")
            sav_cell = ws.cell(row=curr_row, column=5, value=f"=IF(B{curr_row}>0, D{curr_row}/B{curr_row}, 0)")
            cnt_cell = ws.cell(row=curr_row, column=6, value=f"='{sheet_name}'!D4")

            inc_cell.number_format = '"$"#,##0.00'
            exp_cell.number_format = '"$"#,##0.00'
            net_cell.number_format = '"$"#,##0.00;("$"#,##0.00);"-"'
            sav_cell.number_format = '0.0%'
            cnt_cell.number_format = '#,##0'

            for c in range(1, 7):
                ws.cell(row=curr_row, column=c).border = self.thin_border
            curr_row += 1

        # Annual Totals Row
        tot_row = curr_row
        end_breakdown = max(tot_row - 1, start_breakdown)
        ws.cell(row=tot_row, column=1, value="ANNUAL TOTAL").font = self.bold_font
        ws.cell(row=tot_row, column=2, value=f"=SUM(B{start_breakdown}:B{end_breakdown})").font = self.bold_font
        ws.cell(row=tot_row, column=3, value=f"=SUM(C{start_breakdown}:C{end_breakdown})").font = self.bold_font
        ws.cell(row=tot_row, column=4, value=f"=B{tot_row}-C{tot_row}").font = self.bold_font
        ws.cell(row=tot_row, column=5, value=f"=IF(B{tot_row}>0, D{tot_row}/B{tot_row}, 0)").font = self.bold_font
        ws.cell(row=tot_row, column=6, value=f"=SUM(F{start_breakdown}:F{end_breakdown})").font = self.bold_font

        ws.cell(row=tot_row, column=2).number_format = '"$"#,##0.00'
        ws.cell(row=tot_row, column=3).number_format = '"$"#,##0.00'
        ws.cell(row=tot_row, column=4).number_format = '"$"#,##0.00;("$"#,##0.00);"-"'
        ws.cell(row=tot_row, column=5).number_format = '0.0%'
        ws.cell(row=tot_row, column=6).number_format = '#,##0'

        for c in range(1, 7):
            ws.cell(row=tot_row, column=c).border = self.double_bottom_border

        # BLOCK 1: Annual KPI Cards (Row 3-4)
        kpi_headers = ["Annual Income", "Annual Expenses", "Annual Net Cashflow", "Annual Savings Rate"]
        kpi_formulas = [f"=B{tot_row}", f"=C{tot_row}", f"=D{tot_row}", f"=E{tot_row}"]

        for i, (kh, kf) in enumerate(zip(kpi_headers, kpi_formulas), start=1):
            cell_h = ws.cell(row=3, column=i, value=kh)
            cell_v = ws.cell(row=4, column=i, value=kf)

            cell_h.fill = self.soft_blue_fill
            cell_h.font = self.sub_header_font
            cell_v.fill = self.kpi_fill
            cell_v.font = self.bold_font

            if i in [1, 2, 3]:
                cell_v.number_format = '"$"#,##0.00'
            elif i == 4:
                cell_v.number_format = '0.0%'

        # BLOCK 3: Annual Category Spend Matrix
        curr_row += 3
        ws.cell(row=curr_row, column=1, value=f"{year_str} Category Spend Matrix").font = self.section_font
        curr_row += 1

        cat_headers = ["Category", "Total Spent", "% of Annual Expenses", "Monthly Average"]
        for col_idx, ch in enumerate(cat_headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=ch)
            cell.fill = self.navy_fill
            cell.font = self.header_font

        curr_row += 1

        if not self.df.empty and "posted_date" in self.df.columns and "amount" in self.df.columns:
            year_df = self.df[
                (self.df["posted_date"].dt.strftime("%Y-%m").isin(month_sheets))
                & (~self.df.apply(self._is_transfer, axis=1))
                & (self.df["amount"] < 0)
            ]

            if not year_df.empty:
                cat_summary = (
                    year_df.groupby("category")["amount"]
                    .sum()
                    .abs()
                    .reset_index()
                    .sort_values("amount", ascending=False)
                )

                num_months = max(len(month_sheets), 1)

                for _, row in cat_summary.iterrows():
                    ws.cell(row=curr_row, column=1, value=row["category"]).font = self.regular_font
                    
                    amt_cell = ws.cell(row=curr_row, column=2, value=float(row["amount"]))
                    pct_cell = ws.cell(row=curr_row, column=3, value=f"=B{curr_row}/C{tot_row}")
                    avg_cell = ws.cell(row=curr_row, column=4, value=f"=B{curr_row}/{num_months}")

                    amt_cell.number_format = '"$"#,##0.00'
                    pct_cell.number_format = '0.0%'
                    avg_cell.number_format = '"$"#,##0.00'

                    for c in range(1, 5):
                        ws.cell(row=curr_row, column=c).border = self.thin_border
                    curr_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # =========================================================================
    # MASTER OVERVIEW BUILDER
    # =========================================================================
    def _build_master_overview_sheet(self, ws, years_present: list[str], all_month_sheets: list[str]):
        ws.views.sheetView[0].showGridLines = True
        ws["A1"] = "Master Financial Overview"
        ws["A1"].font = self.title_font

        if not years_present:
            ws["A3"] = "No transaction data available."
            return

        # BLOCK 2: YoY Cashflow Summary Table
        curr_row = 7
        ws.cell(row=curr_row, column=1, value="Year-over-Year Cashflow Summary").font = self.section_font
        curr_row += 1

        headers = ["Year", "Total Income", "Total Expenses", "Net Cashflow", "Savings Rate", "YoY Growth"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=h)
            cell.fill = self.navy_fill
            cell.font = self.header_font

        start_yoy = curr_row + 1
        curr_row += 1

        for idx, year in enumerate(years_present):
            y_sheet = f"{year} Overview"
            year_label = self._get_year_status_label(year, all_month_sheets)
            ws.cell(row=curr_row, column=1, value=year_label).font = self.regular_font

            inc_cell = ws.cell(row=curr_row, column=2, value=f"='{y_sheet}'!A4")
            exp_cell = ws.cell(row=curr_row, column=3, value=f"='{y_sheet}'!B4")
            net_cell = ws.cell(row=curr_row, column=4, value=f"='{y_sheet}'!C4")
            sav_cell = ws.cell(row=curr_row, column=5, value=f"='{y_sheet}'!D4")

            if idx == 0:
                yoy_cell = ws.cell(row=curr_row, column=6, value="-")
            else:
                prev_row = curr_row - 1
                yoy_cell = ws.cell(
                    row=curr_row,
                    column=6,
                    value=f"=IF(D{prev_row}<>0, (D{curr_row}-D{prev_row})/ABS(D{prev_row}), 0)",
                )

            inc_cell.number_format = '"$"#,##0.00'
            exp_cell.number_format = '"$"#,##0.00'
            net_cell.number_format = '"$"#,##0.00;("$"#,##0.00);"-"'
            sav_cell.number_format = '0.0%'
            if idx > 0:
                yoy_cell.number_format = '0.0%'

            for c in range(1, 7):
                ws.cell(row=curr_row, column=c).border = self.thin_border
            curr_row += 1

        # Multi-Year Totals Row
        tot_row = curr_row
        end_yoy = max(tot_row - 1, start_yoy)
        ws.cell(row=tot_row, column=1, value="CUMULATIVE TOTAL").font = self.bold_font
        ws.cell(row=tot_row, column=2, value=f"=SUM(B{start_yoy}:B{end_yoy})").font = self.bold_font
        ws.cell(row=tot_row, column=3, value=f"=SUM(C{start_yoy}:C{end_yoy})").font = self.bold_font
        ws.cell(row=tot_row, column=4, value=f"=B{tot_row}-C{tot_row}").font = self.bold_font
        ws.cell(row=tot_row, column=5, value=f"=IF(B{tot_row}>0, D{tot_row}/B{tot_row}, 0)").font = self.bold_font
        ws.cell(row=tot_row, column=6, value="-").font = self.bold_font

        ws.cell(row=tot_row, column=2).number_format = '"$"#,##0.00'
        ws.cell(row=tot_row, column=3).number_format = '"$"#,##0.00'
        ws.cell(row=tot_row, column=4).number_format = '"$"#,##0.00;("$"#,##0.00);"-"'
        ws.cell(row=tot_row, column=5).number_format = '0.0%'

        for c in range(1, 7):
            ws.cell(row=tot_row, column=c).border = self.double_bottom_border

        # BLOCK 1: Multi-Year Macro KPI Banner
        kpi_headers = ["Cumulative Net Savings", "Total Income", "Total Expenses", "Avg Savings Rate", "Period Covered"]
        kpi_formulas = [
            f"=D{tot_row}",
            f"=B{tot_row}",
            f"=C{tot_row}",
            f"=E{tot_row}",
            f'="{all_month_sheets[0]} to {all_month_sheets[-1]}"',
        ]

        for i, (kh, kf) in enumerate(zip(kpi_headers, kpi_formulas), start=1):
            cell_h = ws.cell(row=3, column=i, value=kh)
            cell_v = ws.cell(row=4, column=i, value=kf)

            cell_h.fill = self.soft_blue_fill
            cell_h.font = self.sub_header_font
            cell_v.fill = self.kpi_fill
            cell_v.font = self.bold_font

            if i in [1, 2, 3]:
                cell_v.number_format = '"$"#,##0.00'
            elif i == 4:
                cell_v.number_format = '0.0%'

        # BLOCK 3: Multi-Year Category Spend Matrix
        curr_row += 3
        ws.cell(row=curr_row, column=1, value="Multi-Year Category Spend Matrix").font = self.section_font
        curr_row += 1

        cat_headers = ["Category", "Total Spent", "% of Lifetime Expenses", "Annual Average"]
        for col_idx, ch in enumerate(cat_headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=ch)
            cell.fill = self.navy_fill
            cell.font = self.header_font

        curr_row += 1

        if not self.df.empty and "posted_date" in self.df.columns and "amount" in self.df.columns:
            filtered_df = self.df[
                (self.df["posted_date"].dt.strftime("%Y-%m").isin(all_month_sheets))
                & (~self.df.apply(self._is_transfer, axis=1))
                & (self.df["amount"] < 0)
            ]

            if not filtered_df.empty:
                cat_summary = (
                    filtered_df.groupby("category")["amount"]
                    .sum()
                    .abs()
                    .reset_index()
                    .sort_values("amount", ascending=False)
                )

                num_years = max(len(years_present), 1)

                for _, row in cat_summary.iterrows():
                    ws.cell(row=curr_row, column=1, value=row["category"]).font = self.regular_font
                    
                    amt_cell = ws.cell(row=curr_row, column=2, value=float(row["amount"]))
                    pct_cell = ws.cell(row=curr_row, column=3, value=f"=B{curr_row}/C{tot_row}")
                    avg_cell = ws.cell(row=curr_row, column=4, value=f"=B{curr_row}/{num_years}")

                    amt_cell.number_format = '"$"#,##0.00'
                    pct_cell.number_format = '0.0%'
                    avg_cell.number_format = '"$"#,##0.00'

                    for c in range(1, 5):
                        ws.cell(row=curr_row, column=c).border = self.thin_border
                    curr_row += 1

        # BLOCK 4: Master Internal Transfers & Reconciliation
        curr_row += 2
        ws.cell(row=curr_row, column=1, value="Master Internal Transfers & Reconciliation").font = self.section_font
        curr_row += 1

        trans_headers = ["Transfer Type", "Inflow (Received)", "Outflow (Paid)", "Net Shift", "Status"]
        for col_idx, th in enumerate(trans_headers, start=1):
            cell = ws.cell(row=curr_row, column=col_idx, value=th)
            cell.fill = self.navy_fill
            cell.font = self.header_font

        curr_row += 1

        if not self.df.empty and "posted_date" in self.df.columns and "amount" in self.df.columns:
            transfers_df = self.df[
                (self.df["posted_date"].dt.strftime("%Y-%m").isin(all_month_sheets))
                & (self.df.apply(self._is_transfer, axis=1))
            ]

            target_subcats = ["Credit Card Payment", "Account Transfer", "Loan Proceeds"]

            for subcat in target_subcats:
                sub_df = transfers_df[transfers_df["subcategory"].astype(str).str.lower() == subcat.lower()] if not transfers_df.empty else pd.DataFrame()
                inflow = sub_df[sub_df["amount"] > 0]["amount"].sum() if not sub_df.empty else 0.0
                outflow = sub_df[sub_df["amount"] < 0]["amount"].sum() if not sub_df.empty else 0.0

                ws.cell(row=curr_row, column=1, value=subcat).font = self.regular_font
                
                in_cell = ws.cell(row=curr_row, column=2, value=float(inflow))
                out_cell = ws.cell(row=curr_row, column=3, value=float(outflow))
                net_cell = ws.cell(row=curr_row, column=4, value=f"=B{curr_row}+C{curr_row}")

                in_cell.number_format = '"$"#,##0.00'
                out_cell.number_format = '"$"#,##0.00;("$"#,##0.00);"-"'
                net_cell.number_format = '"$"#,##0.00;("$"#,##0.00);"-"'

                status_cell = ws.cell(row=curr_row, column=5, value=f'=IF(ABS(D{curr_row})<0.01, "Reconciled", "Net Shift")')
                status_cell.font = self.regular_font

                for c in range(1, 6):
                    ws.cell(row=curr_row, column=c).border = self.thin_border
                curr_row += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)